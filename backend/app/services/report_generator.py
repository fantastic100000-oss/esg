"""Word/PDF 보고서 생성"""
import io
import json
from datetime import datetime
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, PageBreak,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import reportlab.rl_config

from app.config import settings

# ESG 컬러 팔레트
COLOR_GREEN = RGBColor(34, 139, 34)
COLOR_BLUE = RGBColor(0, 70, 127)
COLOR_GRAY = RGBColor(100, 100, 100)

RL_GREEN = colors.HexColor("#228B22")
RL_BLUE = colors.HexColor("#00467F")
RL_LIGHT = colors.HexColor("#F0F7F0")


def _add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    run = p.runs[0] if p.runs else p.add_run(text)
    if level == 1:
        run.font.color.rgb = COLOR_BLUE
    elif level == 2:
        run.font.color.rgb = COLOR_GREEN


def _add_kv_table(doc: Document, data: list[tuple[str, str]]):
    """키-값 2열 표 추가"""
    table = doc.add_table(rows=len(data), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (key, val) in enumerate(data):
        table.cell(i, 0).text = key
        table.cell(i, 1).text = str(val) if val is not None else "-"
    doc.add_paragraph()


def generate_word_report(
    analysis: dict,
    my_company: dict,
    competitors: list[dict],
    output_path: str | Path,
) -> Path:
    """Word 형식 ESG 비교분석 보고서 생성"""
    doc = Document()

    # 페이지 여백
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)

    # 표지
    doc.add_paragraph()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("ESG 지속가능경영 비교분석 보고서")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_BLUE

    doc.add_paragraph()
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.add_run(f"식품업계 동종사 벤치마킹 | {datetime.now().strftime('%Y년 %m월')}")

    doc.add_paragraph()
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.add_run(
        f"자사: {my_company['name']} | "
        f"비교 대상: {', '.join(c['name'] for c in competitors)}"
    ).font.color.rgb = COLOR_GRAY

    doc.add_page_break()

    # 목차
    _add_heading(doc, "목차", 1)
    toc_items = [
        "1. 종합 요약 (Executive Summary)",
        "2. 환경(E) 부문 비교분석",
        "3. 사회(S) 부문 비교분석",
        "4. 지배구조(G) 부문 비교분석",
        "5. 업계 ESG 정책/규제 동향",
        "6. 미디어 평판 분석",
        "7. 전략적 제언",
    ]
    for item in toc_items:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()

    # 분석 본문 삽입
    report_text = analysis.get("text", "")
    _insert_report_text(doc, report_text)

    # ESG 지표 비교표 추가
    doc.add_page_break()
    _add_heading(doc, "부록: ESG 핵심 지표 비교표", 1)
    _add_metrics_comparison_table(doc, my_company, competitors)

    output_path = Path(output_path)
    doc.save(output_path)
    return output_path


def _insert_report_text(doc: Document, text: str):
    """마크다운 유사 텍스트를 Word 문서에 삽입"""
    for line in text.split("\n"):
        line = line.rstrip()
        if not line:
            doc.add_paragraph()
        elif line.startswith("# "):
            _add_heading(doc, line[2:], 1)
        elif line.startswith("## "):
            _add_heading(doc, line[3:], 2)
        elif line.startswith("### "):
            _add_heading(doc, line[4:], 3)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("| "):
            # 마크다운 표는 건너뜀 (별도 처리)
            pass
        else:
            doc.add_paragraph(line)


def _add_metrics_comparison_table(
    doc: Document, my_company: dict, competitors: list[dict]
):
    """ESG 핵심 지표 비교표"""
    all_companies = [my_company] + competitors
    headers = ["지표", "단위"] + [c["name"] for c in all_companies]

    metric_labels = [
        ("E", "ghg_scope1", "온실가스 Scope1"),
        ("E", "ghg_scope2", "온실가스 Scope2"),
        ("E", "energy_total", "에너지 사용량"),
        ("E", "renewable_ratio", "재생에너지 비율"),
        ("E", "water_usage", "용수 사용량"),
        ("E", "waste_recycling_ratio", "폐기물 재활용률"),
        ("S", "employee_count", "임직원 수"),
        ("S", "female_ratio", "여성 임직원 비율"),
        ("S", "injury_rate", "산업재해율"),
        ("G", "outside_director_ratio", "사외이사 비율"),
        ("G", "female_director_ratio", "여성이사 비율"),
    ]

    table = doc.add_table(rows=1 + len(metric_labels), cols=len(headers))
    table.style = "Medium Grid 1 Accent 1"

    # 헤더
    for i, h in enumerate(headers):
        table.cell(0, i).text = h

    # 데이터
    for row_i, (cat, key, label) in enumerate(metric_labels, 1):
        table.cell(row_i, 0).text = label
        unit = "-"
        for col_i, company in enumerate(all_companies, 2):
            metrics = company.get("metrics", {})
            cat_data = metrics.get(cat, {})
            item = cat_data.get(key, {})
            if isinstance(item, dict):
                val = item.get("value")
                unit = item.get("unit", "-") or "-"
                table.cell(row_i, col_i).text = str(val) if val is not None else "-"
            else:
                table.cell(row_i, col_i).text = "-"
        table.cell(row_i, 1).text = unit


def generate_pdf_report(
    analysis: dict,
    my_company: dict,
    competitors: list[dict],
    output_path: str | Path,
) -> Path:
    """PDF 형식 ESG 비교분석 보고서 생성"""
    output_path = Path(output_path)
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=2.5 * cm,
        leftMargin=3 * cm,
        topMargin=2.5 * cm,
        bottomMargin=2.5 * cm,
    )

    styles = getSampleStyleSheet()
    style_title = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        fontSize=22,
        textColor=RL_BLUE,
        spaceAfter=12,
        alignment=1,
    )
    style_h1 = ParagraphStyle(
        "CustomH1",
        parent=styles["Heading1"],
        fontSize=14,
        textColor=RL_BLUE,
        spaceBefore=16,
        spaceAfter=8,
    )
    style_h2 = ParagraphStyle(
        "CustomH2",
        parent=styles["Heading2"],
        fontSize=12,
        textColor=RL_GREEN,
        spaceBefore=10,
        spaceAfter=6,
    )
    style_body = ParagraphStyle(
        "CustomBody",
        parent=styles["Normal"],
        fontSize=10,
        leading=16,
        spaceAfter=4,
    )

    story = []

    # 표지
    story.append(Spacer(1, 3 * cm))
    story.append(Paragraph("ESG 지속가능경영 비교분석 보고서", style_title))
    story.append(Spacer(1, 0.5 * cm))
    story.append(Paragraph(
        f"식품업계 동종사 벤치마킹 | {datetime.now().strftime('%Y년 %m월')}",
        ParagraphStyle("sub", parent=styles["Normal"], fontSize=12, alignment=1),
    ))
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph(
        f"자사: {my_company['name']} | 비교: {', '.join(c['name'] for c in competitors)}",
        ParagraphStyle("meta", parent=styles["Normal"], fontSize=10, alignment=1, textColor=colors.gray),
    ))
    story.append(HRFlowable(width="100%", thickness=1, color=RL_BLUE))
    story.append(PageBreak())

    # 본문
    report_text = analysis.get("text", "")
    for line in report_text.split("\n"):
        line = line.strip()
        if not line:
            story.append(Spacer(1, 6))
        elif line.startswith("# "):
            story.append(Paragraph(line[2:], style_h1))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], style_h2))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], ParagraphStyle(
                "h3", parent=styles["Heading3"], fontSize=11, spaceBefore=8
            )))
        elif line.startswith("- ") or line.startswith("* "):
            story.append(Paragraph(f"• {line[2:]}", style_body))
        else:
            story.append(Paragraph(line, style_body))

    doc.build(story)
    return output_path
